from pathlib import Path
from docx import Document
from pptx import Presentation
import json, re, fitz
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'00-engine/parser/output'; OUT.mkdir(parents=True, exist_ok=True)

def latest(pattern, fallback):
    files=sorted(ROOT.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0] if files else ROOT/fallback

def docx_to_blocks(path):
    doc=Document(path); blocks=[]; tables=[]
    for i,p in enumerate(doc.paragraphs,1):
        txt=p.text.strip()
        if txt:
            blocks.append({'index':i,'style':p.style.name if p.style else None,'text':txt})
    for ti,t in enumerate(doc.tables,1):
        rows=[]
        for r in t.rows:
            rows.append([c.text.strip() for c in r.cells])
        tables.append({'index':ti,'rows':len(rows),'cols':max([len(r) for r in rows], default=0),'data':rows})
    return {'path':str(path),'paragraph_count':len(doc.paragraphs),'nonempty_blocks':len(blocks),'tables_count':len(tables),'blocks':blocks,'tables':tables}

def ppt_to_manifest(path):
    prs=Presentation(path); slides=[]
    render_dir=ROOT/'00-engine/parser/ppt_render/slides_png'
    for idx, slide in enumerate(prs.slides,1):
        texts=[s.text.strip() for s in slide.shapes if hasattr(s,'text') and s.text.strip()]
        slides.append({'slide':idx,'text_blocks':texts,'image':str(render_dir/f'slide_{idx:02d}.png')})
    return {'path':str(path),'slides_count':len(prs.slides),'slides':slides}

items={
 'thesis': docx_to_blocks(latest('02-data/thesis/TESIS_HARNO_CSV_MASTER_LOCKED_*.docx','02-data/thesis/TESIS_HARNO_FORMATTED_UNILA.docx')),
 'paparan': docx_to_blocks(latest('02-data/paparan/NASKAH_PAPARAN_KOMPRE_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx','02-data/paparan/NASKAH_PAPARAN_KOMPRE_HARNO_FINAL.docx')),
 'ppt': ppt_to_manifest(ROOT/'02-data/ppt/PPT KOMPRE-HARNO.pptx')
}
for k,v in items.items():
    (OUT/f'{k}_parsed.json').write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
summary={k:{kk:vv for kk,vv in v.items() if kk not in ['blocks','tables','slides']} for k,v in items.items()}
(OUT/'parse_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT/'parse_summary.json')
