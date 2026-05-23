from pathlib import Path
from docx import Document
import json,datetime
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
work=sorted((ROOT/'02-data/paparan').glob('NASKAH_PAPARAN_KOMPRE_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)[0]
doc=Document(work)
fixes={
79:'• Wilcoxon signed-rank: p = 0,003, signifikan secara nonparametrik, tetapi efek sangat kecil.',
81:'Narasi lisan: Uji normalitas menunjukkan selisih skor tidak berdistribusi normal, sehingga Wilcoxon signed-rank digunakan sebagai acuan utama, sementara paired t-test tetap dilaporkan sebagai pembanding. Paired t-test menghasilkan p = 0,157 atau tidak signifikan, sedangkan Wilcoxon menghasilkan p = 0,003 atau signifikan secara nonparametrik. Namun, effect size Cohen’s dz sebesar 0,070 menunjukkan efek sangat kecil. Karena itu, temuan ditafsirkan hati-hati dan tidak cukup untuk menyatakan keunggulan praktis yang kuat. Order effect juga tidak signifikan, sehingga urutan pengujian tidak memengaruhi hasil.',
99:'• Paired t-test tidak signifikan, sedangkan Wilcoxon signifikan secara nonparametrik; efek praktis sangat kecil.',
101:'• Hipotesis alternatif didukung secara nonparametrik, tetapi dukungan praktisnya sangat lemah karena effect size sangat kecil.',
102:'Narasi lisan: Simpulan penelitian disusun sesuai rumusan masalah. Pertama, tingkat usabilitas kedua chatbot relatif setara secara praktis. Kedua, paired t-test tidak menunjukkan perbedaan rata-rata yang signifikan, sedangkan Wilcoxon menunjukkan perbedaan signifikan secara nonparametrik. Namun, ukuran efek sangat kecil sehingga hasilnya harus ditafsirkan hati-hati dan tidak cukup untuk menyatakan keunggulan praktis yang kuat. Ketiga, analisis klaster tidak menunjukkan aspek usabilitas yang menjadi pembeda dominan.'
}
log=[]
for idx,new in fixes.items():
    p=doc.paragraphs[idx-1]
    before=p.text
    for r in p.runs: r.text=''
    if p.runs: p.runs[0].text=new
    else: p.add_run(new)
    log.append({'location':f'P{idx}','before':before,'after':new})
doc.save(work)
out=ROOT/'00-engine/analyzer/output/paparan_narrative_fix_log.json'
out.write_text(json.dumps({'file':str(work),'timestamp':datetime.datetime.now().isoformat(),'changes':log},ensure_ascii=False,indent=2),encoding='utf-8')
print(out)
