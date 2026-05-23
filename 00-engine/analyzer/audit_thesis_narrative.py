from pathlib import Path
from docx import Document
import json,re
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
work=sorted((ROOT/'02-data/thesis').glob('TESIS_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)[0]
doc=Document(work)
items=[]
patterns=[
 ('ABSOLUTE_NO_DIFFERENCE', r'tidak terdapat perbedaan(?! rata-rata).*Wilcoxon|Kedua hasil tersebut berada di atas taraf signifikansi|H0 tidak ditolak dan H1 tidak didukung'),
 ('OLD_XLSX_G5_NARRATIVE', r'dua nilai.*G5|dikoreksi menggunakan median|di luar rentang skala'),
 ('OVERCLAIM_GENZ_SUPERIOR', r'Gen(?:-|\s)?Z.*lebih (?:baik|unggul|tinggi)|keunggulan.*Gen'),
 ('WILCOXON_SIGNIFICANT_CONTEXT', r'Wilcoxon.*0,003|0,003.*Wilcoxon'),
 ('PAIRED_NOT_SIGNIFICANT_CONTEXT', r'paired.*0,157|0,157.*paired'),
 ('EFFECT_SMALL_CONTEXT', r'efek.*sangat kecil|ukuran efek.*sangat kecil|praktis.*hati-hati'),
]
for i,p in enumerate(doc.paragraphs,1):
    txt=p.text.strip()
    if not txt: continue
    for name,pat in patterns:
        if re.search(pat,txt,re.I):
            severity='INFO'
            if name in ['ABSOLUTE_NO_DIFFERENCE','OLD_XLSX_G5_NARRATIVE','OVERCLAIM_GENZ_SUPERIOR']: severity='REVIEW'
            items.append({'location':f'P{i}','category':name,'severity':severity,'text':txt})
for ti,t in enumerate(doc.tables,1):
    for ri,r in enumerate(t.rows,1):
        rowtxt=' | '.join(c.text.strip() for c in r.cells)
        for name,pat in patterns:
            if re.search(pat,rowtxt,re.I):
                severity='INFO'
                if name in ['ABSOLUTE_NO_DIFFERENCE','OLD_XLSX_G5_NARRATIVE','OVERCLAIM_GENZ_SUPERIOR']: severity='REVIEW'
                items.append({'location':f'T{ti}R{ri}','category':name,'severity':severity,'text':rowtxt})
# logical required pair: whenever Wilcoxon p .003 appears, nearby should mention effect small/cautious
contr=[]
for it in items:
    if it['category']=='WILCOXON_SIGNIFICANT_CONTEXT':
        txt=it['text'].lower()
        if not any(k in txt for k in ['efek','hati-hati','praktis','sangat kecil','nonparametrik']):
            contr.append({'location':it['location'],'issue':'Wilcoxon significant mentioned without immediate cautious/effect-size context','text':it['text']})
# absolute statements still problematic
for it in items:
    if it['severity']=='REVIEW': contr.append({'location':it['location'],'issue':it['category'],'text':it['text']})
report={'file':str(work),'items':items,'contradictions':contr,'counts':{'items':len(items),'contradictions':len(contr)}}
out=ROOT/'00-engine/analyzer/output/thesis_narrative_audit.json'
out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Thesis Narrative Audit','','Role: Peneliti/Akademisi Ilmu Komunikasi - Metode Penelitian Kuantitatif',f'File: `{work}`','',f'Contradictions/Review items: `{len(contr)}`','','## Items Requiring Review']
for x in contr:
    md += [f"### {x['location']} — {x['issue']}",x['text'],'']
md += ['','## Informational Hits']
for x in items:
    if x['severity']=='INFO': md.append(f"- {x['location']} — {x['category']}: {x['text'][:260]}")
mdout=ROOT/'00-engine/analyzer/output/thesis_narrative_audit.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
