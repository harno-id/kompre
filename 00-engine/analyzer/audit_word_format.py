from pathlib import Path
from docx import Document
import json,re,zipfile
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
path=ROOT/'02-data/thesis/TESIS_HARNO_CSV_MASTER_UPDATE_WORKING_20260521-092331.docx'
doc=Document(path)
issues=[]
# styles counts
from collections import Counter
styles=Counter(p.style.name if p.style else 'None' for p in doc.paragraphs)
# headings
headings=[]
for i,p in enumerate(doc.paragraphs,1):
    txt=p.text.strip()
    if txt and p.style and 'Heading' in p.style.name:
        headings.append({'p':i,'style':p.style.name,'text':txt})
# caption audit
captions=[]
for i,p in enumerate(doc.paragraphs,1):
    txt=p.text.strip()
    sty=p.style.name if p.style else ''
    if 'caption' in sty.lower() or re.match(r'^(Tabel|Gambar)\s+',txt,re.I):
        captions.append({'p':i,'style':sty,'text':txt})
        if txt.startswith('Tabel') and sty!='Caption Tabel':
            issues.append({'severity':'WARN','category':'caption_style','location':f'P{i}','message':f'Tabel caption style bukan Caption Tabel: {sty}','text':txt})
# table structure
wide_tables=[]; empty_cells=[]
for ti,t in enumerate(doc.tables,1):
    rows=len(t.rows); cols=max((len(r.cells) for r in t.rows), default=0)
    if cols>6: wide_tables.append({'table':ti,'rows':rows,'cols':cols})
    empt=0
    for r in t.rows:
        for c in r.cells:
            if not c.text.strip(): empt+=1
    if empt and rows*cols>0 and empt/(rows*cols)>0.4:
        empty_cells.append({'table':ti,'empty_cells':empt,'total':rows*cols})
# toc/list styles
for i,p in enumerate(doc.paragraphs,1):
    txt=p.text.strip(); sty=p.style.name if p.style else ''
    if sty.startswith('toc') and not re.search(r'\d+$',txt):
        issues.append({'severity':'WARN','category':'toc','location':f'P{i}','message':'TOC entry tidak berakhir nomor halaman','text':txt})
# required headings
required=['ABSTRAK','ABSTRACT','DAFTAR ISI','DAFTAR TABEL','DAFTAR GAMBAR','I. PENDAHULUAN','II. TINJAUAN PUSTAKA','III. METODE PENELITIAN','IV. HASIL DAN PEMBAHASAN','V. SIMPULAN DAN SARAN','DAFTAR PUSTAKA','LAMPIRAN']
alltext='\n'.join(p.text.strip() for p in doc.paragraphs)
for r in required:
    if r not in alltext:
        issues.append({'severity':'FAIL','category':'required_heading','location':'document','message':f'Heading wajib tidak ditemukan: {r}'})
# sections
sections=[]
for si,s in enumerate(doc.sections,1):
    sections.append({'section':si,'top_cm':round(s.top_margin.cm,2),'bottom_cm':round(s.bottom_margin.cm,2),'left_cm':round(s.left_margin.cm,2),'right_cm':round(s.right_margin.cm,2),'page_width_cm':round(s.page_width.cm,2),'page_height_cm':round(s.page_height.cm,2),'header_cm':round(s.header_distance.cm,2),'footer_cm':round(s.footer_distance.cm,2)})
# field codes quick xml scan
with zipfile.ZipFile(path) as z:
    xml=z.read('word/document.xml').decode('utf-8',errors='ignore')
field_count=xml.count('w:fldChar')+xml.count('w:instrText')
toc_field='TOC' in xml
# residual weird chars / replacement char
weird=[]
for i,p in enumerate(doc.paragraphs,1):
    if '�' in p.text:
        weird.append({'p':i,'text':p.text})
        issues.append({'severity':'WARN','category':'encoding','location':f'P{i}','message':'Replacement character ditemukan','text':p.text[:200]})
report={'file':str(path),'paragraph_count':len(doc.paragraphs),'table_count':len(doc.tables),'section_count':len(doc.sections),'styles':styles.most_common(30),'headings_count':len(headings),'headings':headings,'captions_count':len(captions),'captions':captions,'wide_tables':wide_tables,'empty_cell_tables':empty_cells,'sections':sections,'field_count':field_count,'toc_field_present':toc_field,'issues':issues,'summary':{'fail':sum(1 for x in issues if x['severity']=='FAIL'),'warn':sum(1 for x in issues if x['severity']=='WARN'),'info':sum(1 for x in issues if x['severity']=='INFO')}}
out=ROOT/'00-engine/analyzer/output/word_format_audit.json'
out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Word Format Audit','',f'File: `{path}`','', '## Summary',f"- Paragraphs: `{len(doc.paragraphs)}`",f"- Tables: `{len(doc.tables)}`",f"- Sections: `{len(doc.sections)}`",f"- Captions: `{len(captions)}`",f"- Field count markers: `{field_count}`",f"- TOC field present: `{toc_field}`",f"- FAIL: `{report['summary']['fail']}`",f"- WARN: `{report['summary']['warn']}`",'', '## Issues']
if issues:
    for x in issues: md.append(f"- **{x['severity']}** `{x['category']}` {x['location']}: {x['message']}" + (f" — {x.get('text','')[:180]}" if x.get('text') else ''))
else: md.append('- Tidak ada issue struktural besar terdeteksi.')
md += ['', '## Sections']
for s in sections: md.append(f"- Section {s['section']}: margin L/R/T/B `{s['left_cm']}/{s['right_cm']}/{s['top_cm']}/{s['bottom_cm']}` cm, page `{s['page_width_cm']}x{s['page_height_cm']}` cm")
md += ['', '## Caption Styles']
for c in captions: md.append(f"- P{c['p']} `{c['style']}` — {c['text']}")
md += ['', '## Wide Tables']
for t in wide_tables: md.append(f"- Table {t['table']}: rows `{t['rows']}`, cols `{t['cols']}`")
mdout=ROOT/'00-engine/analyzer/output/word_format_audit.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
print(json.dumps(report['summary'],ensure_ascii=False))
