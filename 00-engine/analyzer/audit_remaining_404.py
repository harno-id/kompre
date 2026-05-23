from pathlib import Path
from docx import Document
import json,re
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
work=sorted((ROOT/'02-data/thesis').glob('TESIS_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)[0]
doc=Document(work)
hits=[]
for i,p in enumerate(doc.paragraphs,1):
    if '404' in p.text:
        hits.append({'location':f'P{i}','type':'paragraph','text':p.text})
for ti,t in enumerate(doc.tables,1):
    for ri,row in enumerate(t.rows,1):
        for ci,cell in enumerate(row.cells,1):
            if '404' in cell.text:
                hits.append({'location':f'T{ti}R{ri}C{ci}','type':'table_cell','text':cell.text})
classified=[]
for h in hits:
    txt=h['text']
    if re.search(r'df\s*=\s*404|df/N.*404|t\(404\)', txt, re.I):
        status='KEEP_DF_404'
        reason='df=404 valid karena N=405 pada paired t-test.'
    elif re.search(r'N\s*nonzero\s*=\s*404|N\s*=\s*404|404\s*responden|404\s*entri|404\s*orang|N=404|N = 404', txt, re.I):
        status='PATCH_TO_405'
        reason='Mengacu jumlah responden/data, harus CSV master N=405.'
    else:
        status='REVIEW_CONTEXT'
        reason='Perlu lihat konteks manual.'
    classified.append({**h,'status':status,'reason':reason})
out=ROOT/'00-engine/analyzer/output/audit_404_remaining.json'
out.write_text(json.dumps({'file':str(work),'count':len(classified),'items':classified},ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Audit Remaining 404','',f'File: `{work}`',f'Count: `{len(classified)}`','']
for x in classified:
    md += [f"## {x['location']} — {x['status']}",f"Reason: {x['reason']}",'',x['text'],'']
mdout=ROOT/'00-engine/analyzer/output/audit_404_remaining.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
for x in classified: print(x['location'], x['status'], x['text'][:300].replace('\n',' | '))
