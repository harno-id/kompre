from pathlib import Path
from docx import Document
import json, datetime
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
work=sorted((ROOT/'02-data/thesis').glob('TESIS_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)[0]
doc=Document(work)
log=[]
# P915 fix
p=doc.paragraphs[914]
before=p.text
after=('2. Hasil pengujian menunjukkan bahwa paired sample t-test tidak menemukan perbedaan rata-rata yang signifikan '
       'antara tingkat usabilitas interaksi komunikasi chatbot bergaya bahasa formal dan chatbot bergaya bahasa Generasi Z '
       '(Sig. = 0,157). Namun, karena distribusi selisih skor tidak normal, uji Wilcoxon signed-rank digunakan sebagai acuan '
       'nonparametrik dan menunjukkan hasil signifikan (Sig. = 0,003). Dengan demikian, terdapat indikasi perbedaan secara '
       'nonparametrik, tetapi ukuran efek Cohen’s dz sebesar 0,070 menunjukkan bahwa perbedaan tersebut sangat kecil secara '
       'praktis. Oleh karena itu, temuan ini perlu ditafsirkan secara hati-hati dan tidak cukup untuk menyatakan keunggulan '
       'praktis yang kuat pada salah satu gaya bahasa chatbot.')
for r in p.runs: r.text=''
if p.runs: p.runs[0].text=after
else: p.add_run(after)
log.append({'location':'P915','before':before,'after':after})
# Table 21 row 3 fix (data cleaning table)
t=doc.tables[20]
row=t.rows[2]
before=[c.text for c in row.cells]
if len(row.cells)>=2:
    row.cells[0].text='Nilai di luar rentang skala 1–5'
    row.cells[1].text='Tidak ditemukan pada CSV raw master; seluruh nilai CUQ berada dalam rentang valid 1–5'
after=[c.text for c in row.cells]
log.append({'location':'T21R3','before':before,'after':after})
doc.save(work)
out=ROOT/'00-engine/analyzer/output/final_narrative_patch_log.json'
out.write_text(json.dumps({'file':str(work),'timestamp':datetime.datetime.now().isoformat(),'changes':log},ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Final Narrative Patch Log','',f'File: `{work}`','']
for x in log:
    md += [f"## {x['location']}",'**Before:**',str(x['before']),'','**After:**',str(x['after']),'']
mdout=ROOT/'00-engine/analyzer/output/final_narrative_patch_log.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
