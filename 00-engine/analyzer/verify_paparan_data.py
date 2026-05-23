from pathlib import Path
from docx import Document
import json,re
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
work=sorted((ROOT/'02-data/paparan').glob('NASKAH_PAPARAN_KOMPRE_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)[0]
doc=Document(work)
texts=[]
for i,p in enumerate(doc.paragraphs,1): texts.append((f'P{i}',p.text))
for ti,t in enumerate(doc.tables,1):
    for ri,r in enumerate(t.rows,1):
        for ci,c in enumerate(r.cells,1): texts.append((f'T{ti}R{ri}C{ci}',c.text))
expected=['405','72,92','73,50','0,57','0,157','0,003','0,070','0,905','0,922']
old=['54,59','54,70','0,313','0,091','0,866','0,890','0,050','404 responden','404 entri']
results=[]
for v in expected:
    loc=[l for l,t in texts if v in t]
    results.append({'value':v,'mode':'expected','count':len(loc),'locations':loc[:20],'status':'PASS' if loc else 'WARN_MISSING'})
for v in old:
    loc=[l for l,t in texts if v in t]
    results.append({'value':v,'mode':'old_absent','count':len(loc),'locations':loc[:20],'status':'PASS' if not loc else 'FAIL_OLD_PRESENT'})
# narrative risk
risk=[]
for loc,txt in texts:
    if re.search(r'tidak terdapat perbedaan signifikan|H0 tidak ditolak|Wilcoxon.*0,003',txt,re.I):
        risk.append({'location':loc,'text':txt})
report={'file':str(work),'results':results,'risk_hits':risk,'summary':{'fail_old':sum(1 for r in results if r['status']=='FAIL_OLD_PRESENT'),'warn_missing':sum(1 for r in results if r['status']=='WARN_MISSING'),'risk_hits':len(risk)}}
out=ROOT/'00-engine/analyzer/output/paparan_data_verify.json'
out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Paparan Data Verify','',f'File: `{work}`','',f"- FAIL old present: `{report['summary']['fail_old']}`",f"- WARN expected missing: `{report['summary']['warn_missing']}`",f"- Risk hits: `{report['summary']['risk_hits']}`",'', '## Value checks']
for r in results: md.append(f"- **{r['status']}** {r['mode']} `{r['value']}` count `{r['count']}` loc `{', '.join(r['locations'])}`")
md += ['', '## Risk hits']
for x in risk: md.append(f"- {x['location']}: {x['text'][:500]}")
mdout=ROOT/'00-engine/analyzer/output/paparan_data_verify.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
print(json.dumps(report['summary'],ensure_ascii=False))
