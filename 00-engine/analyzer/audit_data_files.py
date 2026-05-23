from pathlib import Path
from collections import Counter
import json, re, statistics
from docx import Document
from pptx import Presentation
import openpyxl

root = Path(r'd:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
out = root / '00-engine' / 'analyzer' / 'data_audit_report.json'
files = {
    'thesis': root/'02-data'/'thesis'/'TESIS_HARNO_FORMATTED_UNILA.docx',
    'paparan': root/'02-data'/'paparan'/'NASKAH_PAPARAN_KOMPRE_HARNO_FINAL.docx',
    'ppt': root/'02-data'/'ppt'/'PPT KOMPRE-HARNO.pptx',
    'cuq': root/'02-data'/'cuq'/'CUQ Data mentah.xlsx',
}

def text_stats(texts):
    joined='\n'.join(t for t in texts if t)
    words=re.findall(r"\b\w+\b", joined.lower(), flags=re.UNICODE)
    return {
        'paragraphs_or_blocks': len(texts),
        'nonempty_blocks': sum(1 for t in texts if t.strip()),
        'chars': len(joined),
        'words': len(words),
        'top_terms': Counter(w for w in words if len(w)>4).most_common(20)
    }

def audit_docx(path):
    doc=Document(path)
    paras=[p.text.strip() for p in doc.paragraphs]
    headings=[]
    styles=Counter()
    for p in doc.paragraphs:
        styles[p.style.name if p.style else 'None'] += 1
        if p.style and ('Heading' in p.style.name or 'Judul' in p.style.name):
            if p.text.strip(): headings.append({'style':p.style.name,'text':p.text.strip()[:250]})
    tables=[]
    for i,t in enumerate(doc.tables,1):
        rows=len(t.rows); cols=max((len(r.cells) for r in t.rows), default=0)
        sample=[]
        for r in t.rows[:3]: sample.append([c.text.strip()[:80] for c in r.cells[:5]])
        tables.append({'index':i,'rows':rows,'cols':cols,'sample':sample})
    return {'path':str(path),'size':path.stat().st_size,'stats':text_stats(paras),'styles':styles.most_common(30),'headings':headings[:80],'tables_count':len(tables),'tables':tables[:20]}

def audit_pptx(path):
    prs=Presentation(path)
    slides=[]; all_text=[]
    for idx, slide in enumerate(prs.slides,1):
        texts=[]
        for shape in slide.shapes:
            if hasattr(shape,'text') and shape.text.strip(): texts.append(shape.text.strip())
        all_text.extend(texts)
        slides.append({'slide':idx,'text_blocks':len(texts),'texts':[t[:300] for t in texts]})
    return {'path':str(path),'size':path.stat().st_size,'slides_count':len(prs.slides),'stats':text_stats(all_text),'slides':slides}

def audit_xlsx(path):
    wb=openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets=[]
    for ws in wb.worksheets:
        rows=list(ws.iter_rows(values_only=True))
        nonempty=sum(1 for r in rows if any(v is not None for v in r))
        max_cols=max((len(r) for r in rows), default=0)
        headers=[str(v) if v is not None else '' for v in (rows[0] if rows else [])]
        sample=[list(r[:10]) for r in rows[:5]]
        # numeric columns basic
        num_cols=[]
        for c in range(max_cols):
            vals=[]
            for r in rows[1:]:
                if c < len(r) and isinstance(r[c], (int,float)):
                    vals.append(float(r[c]))
            if vals:
                num_cols.append({'col':c+1,'header':headers[c] if c<len(headers) else '', 'n':len(vals), 'min':min(vals), 'max':max(vals), 'mean':round(sum(vals)/len(vals),4)})
        sheets.append({'name':ws.title,'rows':len(rows),'nonempty_rows':nonempty,'max_cols':max_cols,'headers':headers,'sample':sample,'numeric_columns':num_cols[:80]})
    return {'path':str(path),'size':path.stat().st_size,'sheets_count':len(sheets),'sheets':sheets}

report={'files':{}}
report['files']['thesis']=audit_docx(files['thesis'])
report['files']['paparan']=audit_docx(files['paparan'])
report['files']['ppt']=audit_pptx(files['ppt'])
report['files']['cuq']=audit_xlsx(files['cuq'])
out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(out)
