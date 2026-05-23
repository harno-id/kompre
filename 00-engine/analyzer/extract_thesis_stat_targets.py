from pathlib import Path
from docx import Document
import re,json
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
files=sorted((ROOT/'02-data/thesis').glob('TESIS_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)
work=files[0]
doc=Document(work)
patterns=[r'404',r'54[,.]',r'0[,.]866',r'0[,.]890',r'1[,.]010',r'0[,.]313',r'0[,.]091',r'Cohen',r'Wilcoxon',r'paired',r'Paired',r'reliabilitas',r'validitas',r'normalitas',r'CUQ']
rx=re.compile('|'.join(patterns), re.I)
items=[]
for i,p in enumerate(doc.paragraphs,1):
    txt=p.text.strip()
    if txt and rx.search(txt):
        items.append({'type':'paragraph','index':i,'style':p.style.name if p.style else None,'text':txt})
for ti,t in enumerate(doc.tables,1):
    text=' | '.join(c.text.strip() for r in t.rows for c in r.cells)
    if rx.search(text):
        sample=[]
        for r in t.rows[:8]: sample.append([c.text.strip() for c in r.cells[:8]])
        items.append({'type':'table','index':ti,'rows':len(t.rows),'cols':max((len(r.cells) for r in t.rows), default=0),'sample':sample})
out=ROOT/'00-engine/analyzer/output/thesis_stat_targets.json'
out.write_text(json.dumps({'working_copy':str(work),'matches':items},ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Thesis Statistic Targets','',f'Working copy: `{work}`','',f'Matches: `{len(items)}`','']
for it in items:
    if it['type']=='paragraph': md.append(f"- P{it['index']} ({it['style']}): {it['text'][:400]}")
    else: md.append(f"- Table {it['index']} rows={it['rows']} cols={it['cols']} sample={it['sample'][:2]}")
mdout=ROOT/'00-engine/analyzer/output/thesis_stat_targets.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
