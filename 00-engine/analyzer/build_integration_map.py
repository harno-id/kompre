from pathlib import Path
from docx import Document
import json,re
root=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
# thesis headings
th=Document(root/'02-data/thesis/TESIS_HARNO_FORMATTED_UNILA.docx')
headings=[]
for p in th.paragraphs:
    txt=p.text.strip()
    if txt and p.style and ('Heading' in p.style.name):
        headings.append({'style':p.style.name,'text':txt})
# paparan slides
pap=Document(root/'02-data/paparan/NASKAH_PAPARAN_KOMPRE_HARNO_FINAL.docx')
slides=[]
for p in pap.paragraphs:
    txt=p.text.strip()
    m=re.match(r'^Slide\s+(\d+)\s+[—-]\s+(.+)$', txt)
    if m: slides.append({'slide':int(m.group(1)),'title':m.group(2)})
# mapping heuristic
map_rules=[
 (range(1,2),'Cover/Identitas',['ABSTRAK']),
 (range(2,5),'Pendahuluan',['I. PENDAHULUAN','1.1 \tLatar Belakang','1.2 \tRumusan Masalah','1.3 \tTujuan Penelitian']),
 (range(5,7),'Konsep dan Teori',['II. TINJAUAN PUSTAKA','2.2 \tLandasan Teoretis','2.3 \tChatbot sebagai Media Komunikasi Institusional','2.4 \tGaya Bahasa Chatbot dan Generasi Z','2.5 \tUsabilitas Interaksi Komunikasi dan CUQ']),
 (range(7,9),'Metode',['III. METODE PENELITIAN','3.1 \tJenis dan Desain Penelitian','3.5 \tPopulasi, Sampel, dan Teknik Penentuan Sampel','3.6 \tInstrumen Penelitian','3.10 \tTeknik Analisis Data']),
 (range(9,12),'Hasil',['IV. HASIL DAN PEMBAHASAN','4.2 Deskripsi Data Penelitian','4.3 Hasil Uji Instrumen Penelitian','4.5 Hasil Pengujian Hipotesis']),
 (range(12,14),'Pembahasan/Implikasi',['4.6 Pembahasan','5.2 Implikasi Teoretis','5.3 Implikasi Praktis dan Rekomendasi Strategis']),
 (range(14,16),'Simpulan/Saran',['V. SIMPULAN DAN SARAN','5.1 Simpulan','5.4 Keterbatasan Penelitian','5.5 Saran Penelitian Lanjutan']),
]
integration=[]
for s in slides:
    bucket=''; refs=[]
    for rr,b,r in map_rules:
        if s['slide'] in rr: bucket=b; refs=r
    integration.append({'slide':s['slide'],'slide_title':s['title'],'section_bucket':bucket,'thesis_refs':refs,'ppt_image':str(root/f"00-engine/parser/ppt_render/slides_png/slide_{s['slide']:02d}.png")})
report={'thesis_heading_count':len(headings),'paparan_slide_count':len(slides),'ppt_image_count':15,'integration_map':integration,'notes':['PPT text layer minimal; use rendered PNG for visual/OCR review.','Paparan has slide titles 1-15 and can serve as semantic text source for PPT.']}
out=root/'00-engine/analyzer/integration_map_report.json'
out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(out)
for x in integration: print(f"{x['slide']:02d} | {x['section_bucket']} | {x['slide_title']}")
