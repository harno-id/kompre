from pathlib import Path
from docx import Document
import re,json
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
work=sorted((ROOT/'02-data/thesis').glob('TESIS_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)[0]
doc=Document(work)
checks={
 'old_404':0,'old_54_59':0,'old_54_70':0,'old_0_313':0,'old_0_091':0,'old_alpha_0866':0,'old_alpha_0890':0,
 'new_405':0,'new_72_92':0,'new_73_50':0,'new_0_157':0,'new_0_003':0,'new_alpha_0905':0,'new_alpha_0922':0
}
patterns={'old_404':'404','old_54_59':'54,59','old_54_70':'54,70','old_0_313':'0,313','old_0_091':'0,091','old_alpha_0866':'0,866','old_alpha_0890':'0,890','new_405':'405','new_72_92':'72,92','new_73_50':'73,50','new_0_157':'0,157','new_0_003':'0,003','new_alpha_0905':'0,905','new_alpha_0922':'0,922'}
alltexts=[]
for p in doc.paragraphs: alltexts.append(p.text)
for t in doc.tables:
    for r in t.rows:
        for c in r.cells: alltexts.append(c.text)
joined='\n'.join(alltexts)
for k,p in patterns.items(): checks[k]=joined.count(p)
# extract patched focus paragraphs
focus=[]
rx=re.compile(r'405|72,92|73,50|0,157|0,003|0,905|0,922|Wilcoxon|paired|Cohen|Cronbach|Shapiro',re.I)
for i,p in enumerate(doc.paragraphs,1):
    if rx.search(p.text): focus.append({'p':i,'text':p.text[:500]})
report={'file':str(work),'checks':checks,'focus_count':len(focus),'focus_first80':focus[:80]}
out=ROOT/'00-engine/analyzer/output/thesis_patch_verify.json'
out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Thesis Patch Verify','',f'File: `{work}`','', '## Counts']+[f'- {k}: `{v}`' for k,v in checks.items()]+['','## Focus paragraphs']+[f"- P{x['p']}: {x['text']}" for x in focus[:80]]
mdout=ROOT/'00-engine/analyzer/output/thesis_patch_verify.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
print(json.dumps(checks,ensure_ascii=False,indent=2))
