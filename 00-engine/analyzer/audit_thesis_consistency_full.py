from pathlib import Path
from docx import Document
import json,re
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
work=sorted((ROOT/'02-data/thesis').glob('TESIS_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)[0]
st=json.loads((ROOT/'00-engine/analyzer/output/cuq_statistics_report.json').read_text(encoding='utf-8'))
doc=Document(work)
texts=[]
for i,p in enumerate(doc.paragraphs,1): texts.append(('P'+str(i),p.text))
for ti,t in enumerate(doc.tables,1):
  for ri,r in enumerate(t.rows,1):
    for ci,c in enumerate(r.cells,1): texts.append((f'T{ti}R{ri}C{ci}',c.text))
joined='\n'.join(t for _,t in texts)
expected={
 'N responden/data':'405',
 'Formal mean':'72,92',
 'Formal SD':'16,95',
 'Gen-Z mean':'73,50',
 'Gen-Z SD':'18,50',
 'Diff mean':'0,57',
 'Paired t':'1,418',
 'Paired df':'404',
 'Paired p':'0,157',
 'Wilcoxon W':'16453,5',
 'Wilcoxon p':'0,003',
 'Cohen dz':'0,070',
 'Formal alpha':'0,905',
 'Gen-Z alpha':'0,922',
 'Shapiro W':'0,836',
}
old_bad=['54,59','54,70','0,313','0,091','0,866','0,890','0,050','0.313','0.091','0.866','0.890','0.050']
results=[]
for label,val in expected.items():
    loc=[loc for loc,txt in texts if val in txt]
    results.append({'metric':label,'expected':val,'count':len(loc),'locations':loc[:30],'status':'PASS' if loc else 'MISSING'})
for val in old_bad:
    loc=[loc for loc,txt in texts if val in txt]
    results.append({'metric':'old_value_should_absent','expected_absent':val,'count':len(loc),'locations':loc[:30],'status':'PASS' if not loc else 'FAIL_OLD_VALUE_PRESENT'})
# semantic table checks
checks=[]
checks.append({'check':'All old major statistics removed','status':'PASS' if all(r['status']=='PASS' for r in results if r['metric']=='old_value_should_absent') else 'FAIL'})
checks.append({'check':'All new major statistics present','status':'PASS' if all(r['status']=='PASS' for r in results if r['metric']!='old_value_should_absent') else 'WARN'})
report={'file':str(work),'results':results,'summary':checks}
out=ROOT/'00-engine/analyzer/output/thesis_consistency_audit.json'
out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Thesis Consistency Audit','','Role: Peneliti/Akademisi Ilmu Komunikasi - Metode Penelitian Kuantitatif','',f'File: `{work}`','','## Summary']
for c in checks: md.append(f"- **{c['status']}** — {c['check']}")
md+=['','## Metric Checks']
for r in results:
    if 'expected_absent' in r:
        md.append(f"- **{r['status']}** old `{r['expected_absent']}` count `{r['count']}` loc `{', '.join(r['locations'][:10])}`")
    else:
        md.append(f"- **{r['status']}** {r['metric']} `{r['expected']}` count `{r['count']}` loc `{', '.join(r['locations'][:10])}`")
mdout=ROOT/'00-engine/analyzer/output/thesis_consistency_audit.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
