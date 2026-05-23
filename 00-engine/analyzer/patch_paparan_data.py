from pathlib import Path
from docx import Document
import json,re,datetime
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
work=sorted((ROOT/'02-data/paparan').glob('NASKAH_PAPARAN_KOMPRE_HARNO_CSV_MASTER_UPDATE_WORKING_*.docx'), key=lambda p:p.stat().st_mtime, reverse=True)[0]
doc=Document(work)
log=[]
repls=[
('404 responden','405 responden'),('404 entri responden','405 entri responden'),('404 orang','405 orang'),('N=404','N=405'),('N = 404','N = 405'),
('54,59','72,92'),('54.59','72.92'),('54,70','73,50'),('54.70','73.50'),('0,11','0,57'),('0.11','0.57'),
('6,59','16,95'),('6.59','16.95'),('7,21','18,50'),('7.21','18.50'),
('0,313','0,157'),('0.313','0.157'),('0,091','0,003'),('0.091','0.003'),
('1,010','1,418'),('1.010','1.418'),('df = 403','df = 404'),('t(403)','t(404)'),
('0,050','0,070'),('0.050','0.070'),('0,866','0,905'),('0.866','0.905'),('0,890','0,922'),('0.890','0.922'),
('dua nilai pada item G5 yang berada di luar rentang skala Likert 1–5','seluruh nilai CUQ pada CSV raw master berada dalam rentang valid skala Likert 1–5'),
('nilai G5 yang berada di luar rentang','nilai CUQ di luar rentang pada CSV raw master'),
]
for i,p in enumerate(doc.paragraphs,1):
    before=p.text; after=before
    for old,new in repls: after=after.replace(old,new)
    # targeted semantic fixes
    after=after.replace('tidak terdapat perbedaan signifikan antara kedua kondisi', 'paired t-test tidak menunjukkan perbedaan signifikan, sedangkan Wilcoxon menunjukkan perbedaan signifikan secara nonparametrik dengan efek sangat kecil')
    after=after.replace('tidak terdapat perbedaan yang signifikan antara kedua kondisi', 'paired t-test tidak menunjukkan perbedaan signifikan, sedangkan Wilcoxon menunjukkan perbedaan signifikan secara nonparametrik dengan efek sangat kecil')
    if after!=before:
        for r in p.runs: r.text=''
        if p.runs: p.runs[0].text=after
        else: p.add_run(after)
        log.append({'location':f'P{i}','before':before,'after':after})
for ti,t in enumerate(doc.tables,1):
    for ri,row in enumerate(t.rows,1):
        for ci,cell in enumerate(row.cells,1):
            before=cell.text; after=before
            for old,new in repls: after=after.replace(old,new)
            if after!=before:
                cell.text=after
                log.append({'location':f'T{ti}R{ri}C{ci}','before':before,'after':after})
doc.save(work)
outdir=ROOT/'00-engine/analyzer/output'
out=(outdir/'paparan_data_patch_log.json')
out.write_text(json.dumps({'file':str(work),'timestamp':datetime.datetime.now().isoformat(),'changes':log,'change_count':len(log)},ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Paparan Data Patch Log','',f'File: `{work}`',f'Change count: `{len(log)}`','']
for x in log:
    md += [f"## {x['location']}",'**Before:**',x['before'],'','**After:**',x['after'],'']
mdout=outdir/'paparan_data_patch_log.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(work)
print(mdout)
print(len(log))
