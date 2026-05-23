from pathlib import Path
from docx import Document
import json, datetime
ROOT=Path(r'D:\DATA\TESIS\KOMPRE\HARNO-KOMPRE')
path=ROOT/'02-data/thesis/TESIS_HARNO_CSV_MASTER_UPDATE_WORKING_20260521-092331.docx'
doc=Document(path)
log=[]
for idx in [516,551,791]:
    p=doc.paragraphs[idx-1]
    before={'p':idx,'style':p.style.name if p.style else None,'text':p.text}
    if not p.text.strip() and p.style and p.style.name in ['Caption Tabel','Tabel Caption']:
        p.style=doc.styles['Normal']
    after={'p':idx,'style':p.style.name if p.style else None,'text':p.text}
    log.append({'location':f'P{idx}','before':before,'after':after})
# ensure P792 Caption Tabel
p=doc.paragraphs[791]
before={'p':792,'style':p.style.name if p.style else None,'text':p.text}
if p.text.strip().startswith('Tabel 4.7') and p.style.name!='Caption Tabel':
    p.style=doc.styles['Caption Tabel']
after={'p':792,'style':p.style.name if p.style else None,'text':p.text}
log.append({'location':'P792','before':before,'after':after})
doc.save(path)
out=ROOT/'00-engine/analyzer/output/caption_cleanup_log.json'
out.write_text(json.dumps({'file':str(path),'timestamp':datetime.datetime.now().isoformat(),'changes':log},ensure_ascii=False,indent=2),encoding='utf-8')
md=['# Caption Cleanup Log','',f'File: `{path}`','']
for x in log: md += [f"## {x['location']}",f"Before: `{x['before']}`",f"After: `{x['after']}`",'']
mdout=ROOT/'00-engine/analyzer/output/caption_cleanup_log.md'
mdout.write_text('\n'.join(md),encoding='utf-8')
print(mdout)
